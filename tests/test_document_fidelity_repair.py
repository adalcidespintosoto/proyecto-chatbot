"""Tests for bounded same-model repair of source-to-Markdown fidelity issues."""

from __future__ import annotations

import hashlib
import json
import base64
import io
import sys
from pathlib import Path
from types import SimpleNamespace

import pytest

PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from app.services.bot_instance_service import BotRAGRuntime
from app.services.knowledge_base_service import audit_source_against_markdown


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _prepare_runtime(monkeypatch, tmp_path: Path, audit_reports, repairs):
    source = tmp_path / "manual.pptx"
    source.write_bytes(b"pptx-test-source")
    output = tmp_path / "sanitized"
    markdown = output / "markdown" / "manual.md"
    markdown.parent.mkdir(parents=True)
    markdown.write_text("# Manual\nContenido inicial. (Fuente: diapositiva 1)\n", encoding="utf-8")
    record = {
        "status": "success",
        "source_sha256": _sha256(source),
        "markdown_relative_path": "markdown/manual.md",
        "markdown_sha256": _sha256(markdown),
        "source_markers": [["DIAPOSITIVA", 1]],
        "transform_provider": "openai",
        "transform_model": "gpt-6-luna",
    }
    manifest = {"schema_version": 2, "records": {source.name: dict(record)}}
    (output / "manifest.json").write_text(json.dumps(manifest), encoding="utf-8")

    def prepare_markdown(**kwargs):
        kwargs["source_texts"][source.name] = "Archivo fuente: manual.pptx\n[DIAPOSITIVA 1]\nTexto original de la diapositiva."
        kwargs["source_unit_counts"][source.name] = 1
        kwargs["source_visual_slide_numbers"][source.name] = [1]
        return {source.name: dict(record)}, []

    monkeypatch.setattr("scripts.ingest_sanitized_documents._prepare_markdown", prepare_markdown)

    class FakeClient:
        provider = "openai"
        active_model = "gpt-6-luna"

        def __init__(self, provider=None, model=None):
            self.provider = provider
            self.active_model = model
            self.calls = []

        async def chat_completion(self, messages, **kwargs):
            self.calls.append((messages, kwargs))
            return {"content": repairs.pop(0), "provider": self.provider, "model": self.active_model}

    clients = []

    def make_client(**kwargs):
        client = FakeClient(**kwargs)
        clients.append(client)
        return client

    monkeypatch.setattr("app.services.llm_client.LLMClient", make_client)

    reports = list(audit_reports)
    audit_calls = []

    async def audit_source(*args, **kwargs):
        assert kwargs["llm_client"] is clients[0]
        assert kwargs["extracted_source_text"].startswith("Archivo fuente:")
        audit_calls.append(kwargs)
        return reports.pop(0)

    monkeypatch.setattr("app.services.knowledge_base_service.audit_source_against_markdown", audit_source)
    runtime = object.__new__(BotRAGRuntime)
    return runtime, source, output, markdown, clients, audit_calls


def _issue_report(status="revisar"):
    return {
        "status": status,
        "summary": "Falta un dato del documento original.",
        "issues": [{
            "type": "omission",
            "severity": "high",
            "location": "diapositiva 1",
            "source_evidence": "El documento incluye el dato de contacto.",
            "markdown_evidence": "No aparece en el Markdown.",
            "explanation": "Incluir la información omitida sin añadir datos externos.",
        }],
    }


def test_fidelity_issue_is_repaired_with_same_model_then_reaudited(monkeypatch, tmp_path):
    corrected = "# Manual\nContacto: soporte@example.edu. (Fuente: diapositiva 1)\n"
    runtime, source, output, markdown, clients, _ = _prepare_runtime(
        monkeypatch,
        tmp_path,
        [_issue_report(), {"status": "coincide", "summary": "Coincide con el original.", "issues": []}],
        [corrected],
    )

    result = runtime.prepare_document_for_audit(source, output)

    assert result["status"] == "success"
    assert result["fidelity_audit"]["status"] == "coincide"
    assert result["fidelity_repair_attempts"] == 1
    assert "soporte@example.edu" in markdown.read_text(encoding="utf-8")
    assert clients[0].provider == "openai"
    assert clients[0].active_model == "gpt-6-luna"
    messages, options = clients[0].calls[0]
    assert "FUENTE ORIGINAL" in messages[1]["content"]
    assert "MARKDOWN ACTUAL" in messages[1]["content"]
    assert options["fallback_to_ollama"] is False
    stored = json.loads((output / "manifest.json").read_text(encoding="utf-8"))["records"][source.name]
    assert stored["markdown_sha256"] == _sha256(markdown)
    assert stored["fidelity_audit"]["status"] == "coincide"


def test_unresolved_fidelity_is_preserved_for_review_and_not_approved(monkeypatch, tmp_path):
    runtime, source, output, markdown, clients, _ = _prepare_runtime(
        monkeypatch,
        tmp_path,
        [_issue_report(), _issue_report(), _issue_report()],
        ["# Manual\nTexto ajustado. (Fuente: diapositiva 1)\n"] * 2,
    )

    result = runtime.prepare_document_for_audit(source, output)

    assert result["status"] == "review_required"
    assert result["fidelity_audit"]["status"] == "revisar"
    assert result["fidelity_repair_attempts"] == 2
    assert markdown.is_file()
    assert json.loads((output / "manifest.json").read_text(encoding="utf-8"))["records"][source.name]["fidelity_audit"]["status"] == "revisar"
    assert len(clients[0].calls) == 2


def test_unverifiable_audit_is_preserved_without_guessing_a_repair(monkeypatch, tmp_path):
    runtime, source, output, markdown, clients, _ = _prepare_runtime(
        monkeypatch,
        tmp_path,
        [{"status": "no_verificable", "summary": "La fuente no se pudo verificar.", "issues": []}],
        [],
    )

    result = runtime.prepare_document_for_audit(source, output)

    assert result["status"] == "review_required"
    assert result["fidelity_audit"]["status"] == "no_verificable"
    assert result["fidelity_repair_attempts"] == 0
    assert markdown.is_file()
    assert clients[0].calls == []


def test_pptx_with_images_always_gets_visual_audit_before_repair(monkeypatch, tmp_path):
    corrected = "# Manual\nTexto y afiche conservados. (Fuente: diapositiva 1)\n"
    runtime, source, output, _, _, audit_calls = _prepare_runtime(
        monkeypatch,
        tmp_path,
        [_issue_report(), _issue_report(), {"status": "coincide", "summary": "Coincide con imágenes originales.", "issues": []}],
        [corrected],
    )
    visual = [{"slide_number": 1, "content_type": "image/jpeg", "image_bytes": b"slide image"}]
    monkeypatch.setattr(BotRAGRuntime, "_pptx_picture_slide_numbers", staticmethod(lambda _source: [1]))
    monkeypatch.setattr(BotRAGRuntime, "_extract_visual_evidence", staticmethod(lambda _source, _report: visual))

    result = runtime.prepare_document_for_audit(source, output)

    assert result["status"] == "success"
    assert len(audit_calls) == 3
    assert audit_calls[0].get("source_visuals") is None
    assert audit_calls[1]["source_visuals"] == visual
    assert audit_calls[2]["source_visuals"] == visual


def test_visual_fidelity_audit_sends_original_slide_images_and_prior_findings(tmp_path):
    source_dir = tmp_path / "source"
    output_dir = tmp_path / "sanitized"
    source_dir.mkdir()
    markdown_dir = output_dir / "markdown"
    markdown_dir.mkdir(parents=True)
    source = source_dir / "manual.pptx"
    source.write_bytes(b"pptx-original")
    markdown = markdown_dir / "manual.md"
    markdown.write_text("# Manual\nPaso visible. (Fuente: diapositiva 12)\n", encoding="utf-8")
    (output_dir / "manifest.json").write_text(json.dumps({
        "records": {
            source.name: {
                "status": "success",
                "source_sha256": _sha256(source),
                "markdown_relative_path": "markdown/manual.md",
                "markdown_sha256": _sha256(markdown),
            }
        }
    }), encoding="utf-8")

    class FakeClient:
        async def chat_completion(self, messages, **kwargs):
            self.messages = messages
            return {"content": '{"status":"coincide","summary":"La imagen confirma el contenido.","issues":[]}', "provider": "openai", "model": "gpt-6-luna"}

    client = FakeClient()
    image_bytes = b"test-slide-jpeg"
    prior = {"status": "no_verificable", "issues": [{"location": "Diapositiva 12", "explanation": "Validar número del paso."}]}
    import asyncio

    report = asyncio.run(audit_source_against_markdown(
        source_dir,
        output_dir,
        source.name,
        force=True,
        llm_client=client,
        extracted_source_text="Archivo fuente: manual.pptx\n[DIAPOSITIVA 12]\nDescripción visual no verificada.",
        extracted_source_units=1,
        source_visuals=[{"slide_number": 12, "image_bytes": image_bytes}],
        previous_audit=prior,
    ))

    user_content = client.messages[1]["content"]
    assert isinstance(user_content, list)
    assert any(item.get("type") == "image_url" and base64.b64encode(image_bytes).decode("ascii") in item["image_url"]["url"] for item in user_content)
    assert any("REVISIÓN PREVIA" in item.get("text", "") for item in user_content if item.get("type") == "text")
    assert report["visual_evidence_slides"] == [12]


def test_visual_evidence_extracts_all_images_from_compound_slide_ranges(tmp_path):
    from PIL import Image
    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    blank_layout = presentation.slide_layouts[6]
    images = []
    for color in ("red", "blue", "green"):
        buffer = io.BytesIO()
        Image.new("RGB", (40, 30), color=color).save(buffer, format="PNG")
        images.append(buffer.getvalue())
    slide_one = presentation.slides.add_slide(blank_layout)
    slide_one.shapes.add_picture(io.BytesIO(images[0]), Inches(0), Inches(0))
    slide_one.shapes.add_picture(io.BytesIO(images[1]), Inches(1), Inches(0))
    slide_two = presentation.slides.add_slide(blank_layout)
    slide_two.shapes.add_picture(io.BytesIO(images[2]), Inches(0), Inches(0))
    source = tmp_path / "visuals.pptx"
    presentation.save(source)

    visuals = BotRAGRuntime._extract_visual_evidence(
        source,
        {"issues": [{"location": "Diapositivas 1 y 2", "source_evidence": "imagen"}]},
    )

    assert [visual["slide_number"] for visual in visuals] == [1, 1, 2]
    assert all(visual["content_type"] == "image/jpeg" for visual in visuals)


def test_repair_receives_the_same_original_images_used_by_visual_audit():
    class FakeClient:
        provider = "openai"
        active_model = "gpt-6-luna"

        async def chat_completion(self, messages, **kwargs):
            self.messages = messages
            self.options = kwargs
            return {"content": "# Manual\nFondo condonable. (Fuente: diapositiva 8)"}

    image_bytes = b"original-slide-image"
    client = FakeClient()
    runtime = object.__new__(BotRAGRuntime)
    corrected, errors, _ = runtime._repair_markdown_for_fidelity(
        markdown="# Manual\nContenido inicial. (Fuente: diapositiva 8)",
        source_text="[DIAPOSITIVA 8] Descripción automática no verificada.",
        report=_issue_report(),
        record={"source_markers": [["DIAPOSITIVA", 8]]},
        llm_client=client,
        source_visuals=[{"slide_number": 8, "content_type": "image/jpeg", "image_bytes": image_bytes}],
    )

    content = client.messages[1]["content"]
    assert isinstance(content, list)
    assert any(item.get("type") == "image_url" and base64.b64encode(image_bytes).decode("ascii") in item["image_url"]["url"] for item in content)
    assert "imagen original" in client.messages[0]["content"]
    repair_prompt = content[0]["text"]
    assert "todos los hechos concretos incluidos en source_evidence" in repair_prompt
    assert "inclúyela literalmente" in repair_prompt
    assert "conserva los números originales" in repair_prompt
    assert errors == []
    assert corrected.endswith("(Fuente: diapositiva 8)\n")
