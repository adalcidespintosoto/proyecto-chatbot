"""
Script para generar el documento formal en formato Microsoft Word (.docx)
de la Solicitud Técnica Definitiva de Infraestructura, Subdominio y Canal WhatsApp para UniMon.
Ajustado con los datos exactos del equipo S3-RESV-PC36, celular móvil para Meta, 
3 sugerencias de subdominio, GLPI de pruebas y sin requerimiento de UPS.
"""

import sys
import os
from pathlib import Path

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn


def set_cell_background(cell, hex_color):
    """Establece el color de fondo de una celda en formato hexadecimal."""
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)


def set_cell_margins(cell, top=120, bottom=120, left=160, right=160):
    """Añade padding interno a las celdas de una tabla en dxa."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)


def set_table_borders(table, color="CBD5E1"):
    """Aplica bordes limpios y discretos a la tabla."""
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="single" w:sz="6" w:space="0" w:color="{color}"/>'
        f'  <w:left w:val="none"/>'
        f'  <w:bottom w:val="single" w:sz="8" w:space="0" w:color="{color}"/>'
        f'  <w:right w:val="none"/>'
        f'  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
        f'  <w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)


def build_docx_solicitud(output_path: Path):
    doc = Document()

    # 1. Configuración de márgenes estándar (2.54 cm / 1 pulgada)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

        # Header institucional
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("UNIVERSIDAD SIMÓN BOLÍVAR  |  DIRECCIÓN DE TECNOLOGÍA E INFORMACIÓN (DITI)")
        hrun.font.name = "Calibri"
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = RGBColor(120, 144, 156)

        # Footer institucional
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        frun = fp.add_run("UniMon — Asistente Virtual de Soporte Técnico USB  •  Documento Oficial de Aprovisionamiento")
        frun.font.name = "Calibri"
        frun.font.size = Pt(8.5)
        frun.font.color.rgb = RGBColor(140, 150, 160)

    # Paleta de colores ejecutiva
    NAVY = RGBColor(27, 54, 93)       # #1B365D Azul Institucional Profundo
    SLATE = RGBColor(70, 90, 120)     # #465A78 Azul Pizarra
    DARK_TEXT = RGBColor(40, 44, 52)  # #282C34 Texto Principal
    ACCENT = RGBColor(0, 122, 204)    # #007ACC Azul Eléctrico de Acento

    # Badge de cabecera
    p_badge = doc.add_paragraph()
    p_badge.paragraph_format.space_before = Pt(0)
    p_badge.paragraph_format.space_after = Pt(4)
    r_badge = p_badge.add_run("SOLICITUD FORMAL DE RECURSOS DE INFRAESTRUCTURA, REDES Y TELECOMUNICACIONES")
    r_badge.font.name = "Calibri"
    r_badge.font.size = Pt(9.5)
    r_badge.font.bold = True
    r_badge.font.color.rgb = ACCENT

    # Título Principal
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(2)
    p_title.paragraph_format.space_after = Pt(4)
    r_title = p_title.add_run("Aprovisionamiento de Red, Subdominio SSL y Línea Móvil para Canal WhatsApp Business")
    r_title.font.name = "Calibri"
    r_title.font.size = Pt(19)
    r_title.font.bold = True
    r_title.font.color.rgb = NAVY

    # Subtítulo
    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(14)
    r_sub = p_sub.add_run("Despliegue Profesional y Enrutamiento Perimetral para el Asistente Virtual UniMon")
    r_sub.font.name = "Calibri"
    r_sub.font.size = Pt(11)
    r_sub.font.color.rgb = SLATE

    # Tabla de Metadatos de Memorando (2 columnas x 3 filas)
    meta_table = doc.add_table(rows=3, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(meta_table, "E2E8F0")

    meta_items = [
        ("PARA:", "Dirección de Tecnología y Sistemas de Información / Infraestructura, Redes y Telecomunicaciones",
         "FECHA:", "29 de septiembre de 2026"),
        ("DE:", "Adalcides Pintos Soto — Equipo de Desarrollo e Inteligencia Artificial",
         "ESTADO:", "Prioridad Alta / Despliegue Piloto Funcional"),
        ("ASUNTO:", "Aprovisionamiento de IP Fija, Subdominio HTTPS y Celular Corporativo WhatsApp",
         "PROYECTO:", "UniMon — Asistente Virtual de Soporte Técnico TI")
    ]

    for row_idx, data in enumerate(meta_items):
        row = meta_table.rows[row_idx]
        col_data = [(data[0], data[1]), (data[2], data[3])]
        for col_idx, (lbl, val) in enumerate(col_data):
            cell = row.cells[col_idx]
            set_cell_background(cell, "F8FAFC")
            set_cell_margins(cell, top=90, bottom=90, left=140, right=140)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            r1 = p.add_run(lbl + " ")
            r1.font.name = "Calibri"
            r1.font.bold = True
            r1.font.size = Pt(9.5)
            r1.font.color.rgb = NAVY
            r2 = p.add_run(val)
            r2.font.name = "Calibri"
            r2.font.size = Pt(9.5)
            r2.font.color.rgb = DARK_TEXT

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # Helper de llamadas / callouts
    def add_callout(text, title="OBJETIVO ESTRATÉGICO:", border_color="007ACC", bg_color="EFF6FF", title_color=ACCENT):
        table = doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = table.cell(0, 0)
        set_cell_background(cell, bg_color)
        set_cell_margins(cell, top=130, bottom=130, left=170, right=170)

        tcPr = cell._tc.get_or_add_tcPr()
        borders = parse_xml(
            f'<w:tcBorders {nsdecls("w")}>'
            f'  <w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/>'
            f'  <w:top w:val="none"/>'
            f'  <w:bottom w:val="none"/>'
            f'  <w:right w:val="none"/>'
            f'</w:tcBorders>'
        )
        tcPr.append(borders)

        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.15
        r_t = p.add_run(title + "\n")
        r_t.font.name = "Calibri"
        r_t.font.bold = True
        r_t.font.color.rgb = title_color
        r_t.font.size = Pt(10)
        r_b = p.add_run(text)
        r_b.font.name = "Calibri"
        r_b.font.size = Pt(9.5)
        r_b.font.color.rgb = DARK_TEXT
        doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # Helper para encabezados de sección
    def add_sec_heading(title, level=1):
        p = doc.add_paragraph()
        p.paragraph_format.keep_with_next = True
        r = p.add_run(title)
        r.font.name = "Calibri"
        r.font.bold = True
        if level == 1:
            p.paragraph_format.space_before = Pt(16)
            p.paragraph_format.space_after = Pt(6)
            r.font.size = Pt(13.5)
            r.font.color.rgb = NAVY
        elif level == 2:
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(4)
            r.font.size = Pt(11.5)
            r.font.color.rgb = SLATE
        else:
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
            r.font.size = Pt(10.5)
            r.font.color.rgb = DARK_TEXT
        return p

    def add_body_p(text, bold_prefix=None, space_after=5):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            rb = p.add_run(bold_prefix)
            rb.font.name = "Calibri"
            rb.font.bold = True
            rb.font.color.rgb = NAVY
            rb.font.size = Pt(10)
        r = p.add_run(text)
        r.font.name = "Calibri"
        r.font.size = Pt(10)
        r.font.color.rgb = DARK_TEXT
        return p

    def add_bullet_p(text, bold_prefix=None):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            rb = p.add_run(bold_prefix)
            rb.font.name = "Calibri"
            rb.font.bold = True
            rb.font.color.rgb = NAVY
            rb.font.size = Pt(9.5)
        r = p.add_run(text)
        r.font.name = "Calibri"
        r.font.size = Pt(9.5)
        r.font.color.rgb = DARK_TEXT
        return p

    # Callout introductorio
    add_callout(
        "Habilitar la atención institucional automatizada de soporte técnico 24/7 a través de WhatsApp oficial y la web institucional "
        "de la Universidad Simón Bolívar. Se solicita el enrutamiento perimetral hacia la estación de trabajo física asignada "
        "(eliminando el uso de localhost:8000), utilizando un número móvil corporativo exclusivo y manteniendo la integración con el "
        "entorno de pruebas de GLPI para la validación integral y segura de esta fase piloto.",
        "PROPÓSITO Y ALCANCE DE LA SOLICITUD:"
    )

    # ==========================================
    # SECCIÓN 1: CONTEXTO Y JUSTIFICACIÓN
    # ==========================================
    add_sec_heading("1. Contexto y Justificación del Proyecto")
    add_body_p(
        "Se ha desarrollado con éxito el asistente virtual UniMon, una solución institucional basada en Inteligencia Artificial "
        "y arquitectura RAG (Retrieval-Augmented Generation) diseñada para brindar soporte técnico automatizado y autogestión de "
        "servicios de TI a estudiantes, profesores y personal administrativo en las sedes de Barranquilla y Cúcuta."
    )
    add_body_p(
        "El sistema ya se encuentra plenamente integrado con la base de conocimiento documental de la Universidad (más de 115 manuales, "
        "instructivos y normativas institucionales indexadas vectorialmente) y cuenta con módulo de radicación y consulta de tickets en "
        "la mesa de ayuda GLPI. Para esta fase de validación y pruebas de campo, se continuará operando contra la instancia de GLPI de pruebas "
        "(https://pruebas.us5.glpi-network.cloud) con el fin de evaluar el comportamiento y exactitud del bot antes de cualquier impacto en producción.",
        bold_prefix="Alcance de Mesa de Ayuda: "
    )
    add_body_p(
        "Dado que la fase actual aprovecha la aceleración por hardware (GPU dedicada) y modelos de visión locales alojados en "
        "la estación de trabajo asignada, se requiere publicar el servicio bajo un subdominio oficial de la institución "
        "mediante enrutamiento perimetral (Reverse Proxy). De esta manera, se brinda atención oficial vía WhatsApp y Web "
        "con los más altos estándares de presentación y ciberseguridad, sin incurrir en compras o aprovisionamiento de costosos servidores GPU en la nube.",
        bold_prefix="Alojamiento en Estación Local con GPU: "
    )

    # ==========================================
    # SECCIÓN 2: DATOS EXACTOS DE LA ESTACIÓN DE TRABAJO
    # ==========================================
    add_sec_heading("2. Especificaciones Técnicas de la Estación de Trabajo Local (Host del Servicio)")
    add_body_p(
        "Con el propósito de facilitar la configuración inmediata en los switches y servidores DHCP de la Universidad, "
        "se detallan a continuación los parámetros de red exactos del computador donde corre el backend:"
    )

    pc_table = doc.add_table(rows=7, cols=2)
    pc_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(pc_table, "CBD5E1")

    pc_specs = [
        ("Nombre del Equipo (Hostname):", "S3-RESV-PC36"),
        ("Sufijo DNS / Dominio Local:", "S3Direccionti.unisimon.edu.co  /  unisimon.edu.co"),
        ("Adaptador de Red Físico:", "Intel(R) Ethernet Connection (19) I219-LM (Cableado Gigabit 1 Gbps)"),
        ("Dirección Física (MAC Address):", "F4-F1-9E-44-D7-6F"),
        ("Dirección IPv4 Actual (DHCP):", "10.0.12.16 (Máscara: 255.255.255.0 / Puerta de Enlace: 10.0.12.254)"),
        ("Servidor DHCP / DNS Local:", "DHCP: 10.0.30.253  |  DNS: 10.0.30.250"),
        ("Puerto Local del Servicio UniMon:", "Puerto TCP 8000 (Protocolo HTTP interno)")
    ]

    for row_idx, (k, v) in enumerate(pc_specs):
        row = pc_table.rows[row_idx]
        set_cell_background(row.cells[0], "F1F5F9")
        set_cell_background(row.cells[1], "FFFFFF")
        set_cell_margins(row.cells[0], top=70, bottom=70, left=120, right=120)
        set_cell_margins(row.cells[1], top=70, bottom=70, left=120, right=120)
        
        p0 = row.cells[0].paragraphs[0]
        p0.paragraph_format.space_after = Pt(0)
        r0 = p0.add_run(k)
        r0.font.name = "Calibri"
        r0.font.bold = True
        r0.font.size = Pt(9.5)
        r0.font.color.rgb = NAVY

        p1 = row.cells[1].paragraphs[0]
        p1.paragraph_format.space_after = Pt(0)
        r1 = p1.add_run(v)
        r1.font.name = "Calibri"
        r1.font.size = Pt(9.5)
        r1.font.color.rgb = DARK_TEXT

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # ==========================================
    # SECCIÓN 3: REQUERIMIENTOS TÉCNICOS DETALLADOS
    # ==========================================
    add_sec_heading("3. Requerimientos Técnicos Solicitados a las Dependencias de TI")

    # Bloque A: Red Local y Subdominio
    add_sec_heading("A. Red Local, Subdominio Institucional y Reverse Proxy (Infraestructura / Redes)", level=2)
    add_bullet_p(
        "Configuración de una reserva estática por dirección MAC en el servidor DHCP corporativo (10.0.30.253) para la tarjeta Intel I219-LM (MAC: F4-F1-9E-44-D7-6F). "
        "Se solicita fijar la IP actual 10.0.12.16 o asignar una IP fija definitiva en el segmento de red correspondiente.",
        bold_prefix="1. Reserva de IP Estática LAN por MAC: "
    )
    add_bullet_p(
        "Creación de un registro DNS (tipo A o CNAME) institucional apuntando a la IP pública del Reverse Proxy perimetral de la Universidad. "
        "Se presentan tres alternativas a consideración según disponibilidad y nomenclatura institucional:\n"
        "   • Opción 1 (Recomendada): asistenteti.unisimon.edu.co\n"
        "   • Opción 2 (Identidad de Marca): unimon.unisimon.edu.co\n"
        "   • Opción 3 (Descriptiva): soporte-ai.unisimon.edu.co\n"
        "   *(O en su defecto, el subdominio bajo *.unisimon.edu.co que determine la DITI según la política vigente)*.",
        bold_prefix="2. Subdominio Oficial en Zona DNS (3 Sugerencias): "
    )
    add_bullet_p(
        "Certificado SSL/TLS oficial institucional (Wildcard *.unisimon.edu.co o emitido por autoridad reconocida). "
        "Condición mandatoria: Meta WhatsApp Cloud API exige estrictamente conexiones HTTPS con TLS 1.2 o superior para la validación y entrega continua del Webhook.",
        bold_prefix="3. Certificado de Seguridad SSL/TLS (HTTPS Obligatorio): "
    )
    add_bullet_p(
        "Configuración en el Reverse Proxy perimetral corporativo (Nginx, IIS, Fortinet o F5) para publicar el servicio eliminando el uso del puerto 8000 hacia el exterior:\n"
        "   • Entrada pública: https://<subdominio_aprobado>:443  (Redirección automática de HTTP 80 a HTTPS 443)\n"
        "   • Destino interno LAN: http://10.0.12.16:8000\n"
        "   • Parámetros Críticos de Proxy (Para soportar inferencia de IA y adjuntos):\n"
        "       - proxy_read_timeout: 60s  (evita cortes de conexión durante consultas semánticas complejas)\n"
        "       - proxy_connect_timeout: 60s\n"
        "       - client_max_body_size: 20M  (permite recepción de capturas de pantalla para tickets de soporte)\n"
        "       - Cabeceras de trazabilidad: Reenvío de X-Forwarded-For y X-Forwarded-Proto para auditoría de IP.",
        bold_prefix="4. Regla de Reverse Proxy y Timeouts de IA: "
    )
    add_bullet_p(
        "• Inbound (Entrada): Permitir tráfico HTTPS (443) proveniente de los rangos de IP de los servidores de Meta hacia el endpoint del webhook.\n"
        "• Outbound (Salida): Permitir salida HTTPS (443) desde la IP 10.0.12.16 hacia graph.facebook.com (WhatsApp API), api.openai.com / Google Gemini y la instancia de pruebas de GLPI (pruebas.us5.glpi-network.cloud).",
        bold_prefix="5. Políticas de Firewall Perimetral: "
    )

    # Bloque B: WhatsApp Business Platform
    add_sec_heading("B. Canal WhatsApp Business Platform (Telecomunicaciones / Meta Business)", level=2)
    add_bullet_p(
        "Asignación de una línea telefónica móvil corporativa exclusiva (número celular con chip/SIM dedicado). "
        "REQUISITO INDISPENSABLE: Debe ser un número MÓVIL (celular), no línea fija ni PBX/conmutador, para que pueda recibir directamente el SMS o llamada automatizada con el código de verificación de 6 dígitos (OTP) de Meta durante el registro sin tropiezos ni demoras.\n"
        "• La línea móvil no debe estar registrada previamente en ninguna cuenta de WhatsApp personal ni Business.",
        bold_prefix="6. Línea Telefónica MÓVIL Corporativa Exclusiva (Con Recepción OTP): "
    )
    add_bullet_p(
        "Vinculación del número celular corporativo dentro del Administrador Comercial verificado (Meta Business Manager) de la Universidad Simón Bolívar y creación de la App con el producto 'WhatsApp Business Platform' en Meta for Developers.",
        bold_prefix="7. Alta en Meta Business Manager Institucional: "
    )
    add_bullet_p(
        "Aprobación del perfil oficial ante Meta:\n"
        "   • Nombre Visible (Display Name): 'UniMon - Soporte TI Universidad Simón Bolívar'\n"
        "   • Categoría: 'Educación / Servicio Técnico de TI'\n"
        "   • Descripción: 'Asistente virtual oficial de soporte técnico de la Universidad Simón Bolívar. Atención 24/7 en Barranquilla y Cúcuta.'\n"
        "   • Avatar / Foto de Perfil: Logo institucional oficial de la Universidad / UniMon.",
        bold_prefix="8. Perfil de Negocio Institucional (Meta Display Name): "
    )
    add_bullet_p(
        "Para incorporar la conexión en el backend, la dependencia encargada debe suministrar al equipo de desarrollo:\n"
        "   • Phone Number ID: Identificador numérico de 15 dígitos asignado al teléfono móvil en Meta.\n"
        "   • WhatsApp Business Account ID (WABA ID): Identificador de la cuenta comercial en Meta.\n"
        "   • Token de Acceso Permanente (System User Token): Token generado para un 'Usuario del Sistema' con permisos:\n"
        "       - whatsapp_business_messaging\n"
        "       - whatsapp_business_management\n"
        "   • URL de Webhook a registrar en Meta: https://<subdominio_aprobado>/api/whatsapp\n"
        "   • Token de Verificación (Verify Token) configurado en el backend: UniMon_USB_Verify_Token_2026\n"
        "   • Campo de suscripción obligatorio en Meta: En el panel de Meta for Developers (sección WhatsApp > Configuración > Webhook), es indispensable marcar la suscripción al campo 'messages'.",
        bold_prefix="9. Credenciales de la API de WhatsApp y Configuración de Webhook: "
    )

    # ==========================================
    # SECCIÓN 4: MATRIZ CONSOLIDADA DE ENTREGABLES
    # ==========================================
    add_sec_heading("4. Matriz Consolidada de Entregables y Responsabilidades")
    add_body_p("A continuación se consolida la lista exacta de componentes requeridos para la puesta en marcha:")

    table_matriz = doc.add_table(rows=8, cols=5)
    table_matriz.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table_matriz, "CBD5E1")

    headers = ["Ítem", "Componente / Entregable", "Área Responsable", "Especificación Técnica Solicitada", "Prioridad"]

    hdr_cells = table_matriz.rows[0].cells
    for i, h_text in enumerate(headers):
        set_cell_background(hdr_cells[i], "1B365D")
        set_cell_margins(hdr_cells[i], top=100, bottom=100, left=100, right=100)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h_text)
        r.font.name = "Calibri"
        r.font.bold = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(255, 255, 255)

    matriz_data = [
        ("1", "Reserva IP Estática LAN", "Infraestructura / Redes", "IP fija para MAC: F4-F1-9E-44-D7-6F (Host: S3-RESV-PC36)", "Alta"),
        ("2", "Subdominio Oficial DNS", "Infraestructura / DNS", "asistenteti.unisimon.edu.co (u opciones 2/3 sugeridas)", "Alta"),
        ("3", "Certificado SSL/TLS", "Seguridad / Infraestructura", "Certificado oficial HTTPS (TLS 1.2+ obligatorio para Meta)", "Alta"),
        ("4", "Reverse Proxy y Timeouts", "Infraestructura / Redes", "Puerto 443 ──▶ 10.0.12.16:8000 (read_timeout: 60s, 20M)", "Alta"),
        ("5", "Línea Móvil (Celular) OTP", "Telecomunicaciones", "Número celular limpio capaz de recibir SMS/Llamada de Meta", "Alta"),
        ("6", "Alta en Meta Business", "Administrador Meta / DITI", "WABA ID + App WhatsApp Business Platform creada", "Alta"),
        ("7", "Token Permanente y IDs", "Administrador Meta / DITI", "System User Token + Phone Number ID entregados a desarrollo", "Alta")
    ]

    for row_idx, row_data in enumerate(matriz_data, start=1):
        row = table_matriz.rows[row_idx]
        bg_color = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, cell_value in enumerate(row_data):
            cell = row.cells[col_idx]
            set_cell_background(cell, bg_color)
            set_cell_margins(cell, top=75, bottom=75, left=90, right=90)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            if col_idx in [0, 4]:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT

            r = p.add_run(cell_value)
            r.font.name = "Calibri"
            r.font.size = Pt(8.5)
            if col_idx == 0:
                r.font.bold = True
                r.font.color.rgb = NAVY
            elif col_idx == 4:
                r.font.bold = True
                r.font.color.rgb = ACCENT
            else:
                r.font.color.rgb = DARK_TEXT

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # ==========================================
    # SECCIÓN 5: COMPROMISOS DEL EQUIPO SOLICITANTE
    # ==========================================
    add_sec_heading("5. Compromisos Técnicos y Blindaje en la Estación de Trabajo")
    add_body_p("En el computador físico S3-RESV-PC36 donde reside el asistente virtual, el equipo solicitante garantiza:")
    add_bullet_p("Se desactivará por completo el estado de suspensión (Sleep), hibernación y ahorro de energía en la tarjeta de red para operar de forma continua.", bold_prefix="Operación Continua: ")
    add_bullet_p("En la BIOS del equipo se configurará el encendido automático tras restablecimiento de energía (AC Power Recovery -> Power On).", bold_prefix="Auto-Encendido tras Interrupción Eléctrica: ")
    add_bullet_p("El puerto 8000 estará blindado en el Firewall de Windows, autorizando únicamente las peticiones internas provenientes de la IP del Reverse Proxy corporativo.", bold_prefix="Aislamiento y Seguridad Local: ")
    add_bullet_p("El backend de UniMon se configurará como un servicio de fondo con inicio automático en el arranque del sistema operativo.", bold_prefix="Demonio Autónomo: ")

    # ==========================================
    # SECCIÓN 6: BLOQUE DE FIRMAS
    # ==========================================
    add_sec_heading("6. Control de Emisión, Aprobación y Recepción")
    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    sig_table = doc.add_table(rows=2, cols=2)
    sig_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(sig_table, "CBD5E1")

    sig_data = [
        ("SOLICITADO POR:", "RECIBIDO Y APROBADO POR:"),
        ("Adalcides Pintos Soto\nEquipo de Desarrollo e Inteligencia Artificial\nProyecto UniMon — Soporte Técnico USB\nUniversidad Simón Bolívar",
         "Dirección de Tecnología y Sistemas de Información (DITI)\nÁrea de Infraestructura, Redes y Telecomunicaciones\nUniversidad Simón Bolívar")
    ]

    for col_idx in range(2):
        cell_hdr = sig_table.rows[0].cells[col_idx]
        set_cell_background(cell_hdr, "1B365D")
        set_cell_margins(cell_hdr, top=80, bottom=80, left=120, right=120)
        p = cell_hdr.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(sig_data[0][col_idx])
        r.font.name = "Calibri"
        r.font.bold = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(255, 255, 255)

        cell_bdy = sig_table.rows[1].cells[col_idx]
        set_cell_background(cell_bdy, "F8FAFC")
        set_cell_margins(cell_bdy, top=140, bottom=140, left=140, right=140)
        p_b = cell_bdy.paragraphs[0]
        p_b.paragraph_format.space_after = Pt(20)
        p_b.paragraph_format.line_spacing = 1.15
        p_b.add_run("\n\n_________________________________________\n").font.color.rgb = SLATE
        r_txt = p_b.add_run(sig_data[1][col_idx])
        r_txt.font.name = "Calibri"
        r_txt.font.size = Pt(9.5)
        r_txt.font.color.rgb = DARK_TEXT

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))
    print(f"Documento Word definitivo generado exitosamente en: {output_path.resolve()}")


if __name__ == "__main__":
    PROJECT_ROOT = Path(__file__).resolve().parent.parent
    target_file = PROJECT_ROOT / "SOLICITUD_TECNICA_INFRAESTRUCTURA_WHATSAPP_UNIMON_DEFINITIVA.docx"
    build_docx_solicitud(target_file)
